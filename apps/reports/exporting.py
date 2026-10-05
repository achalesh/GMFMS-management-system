import csv
import re
from io import BytesIO, StringIO
from xml.sax.saxutils import escape

from django.core.exceptions import ValidationError
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

from apps.idcards.rendering import font_name, serialized_render


def safe_cell(value):
    if isinstance(value, str):
        value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)
        if value.lstrip().startswith(("=", "+", "-", "@")):
            value = "'" + value
    return value


@serialized_render
def pdf_report(title, headers, rows, description):
    if len(rows) > 1500:
        raise ValidationError(
            "PDF reports are limited to 1,500 rows. Narrow filters or choose Excel/CSV."
        )
    font = font_name()
    style = ParagraphStyle("cell", fontName=font, fontSize=7, leading=10, splitLongWords=True)

    def cell(value):
        value = str(value)
        for chunk in re.split(r"([\u0D00-\u0D7F\u200C\u200D]+)", value):
            selected = "CardMalayalam" if any("\u0d00" <= ch <= "\u0d7f" for ch in chunk) else font
            supported = pdfmetrics.getFont(selected).face.charToGlyph
            if any(
                ord(ch) not in supported
                for ch in chunk
                if not ch.isspace() and ch not in "\u200c\u200d"
            ):
                raise ValidationError(
                    "Some characters are not supported by the PDF fonts. Choose Excel or CSV to preserve all text."
                )
        markup = "".join(
            '<font name="CardMalayalam">' + escape(c) + "</font>"
            if any("\u0d00" <= ch <= "\u0d7f" for ch in c)
            else escape(c)
            for c in re.split(r"([\u0D00-\u0D7F\u200C\u200D]+)", value)
        )
        return Paragraph(
            markup,
            ParagraphStyle(
                "row", parent=style, shaping=any("\u0d00" <= ch <= "\u0d7f" for ch in value)
            ),
        )

    out = BytesIO()
    doc = SimpleDocTemplate(
        out, pagesize=landscape(A4), rightMargin=24, leftMargin=24, topMargin=30, bottomMargin=34
    )
    weights = [
        1.5
        if h in {"Facilitator ID", "Application number", "Name", "District", "Panchayat"}
        else 1.05
        for h in headers
    ]
    table = LongTable(
        [[cell(h) for h in headers]] + [[cell(v) for v in row] for row in rows],
        repeatRows=1,
        colWidths=[doc.width * w / sum(weights) for w in weights],
        hAlign="LEFT",
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#d5eee7")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f7f6")]),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                ("LINEBELOW", (0, 0), (-1, 0), 1, colors.HexColor("#117d69")),
            ]
        )
    )

    def footer(canvas, doc):
        canvas.setFont(font, 7)
        canvas.drawString(
            24,
            18,
            "Gramaswaraj | Administrative report | "
            + timezone.localtime().strftime("%Y-%m-%d %H:%M %Z"),
        )
        canvas.drawRightString(doc.pagesize[0] - 24, 18, f"Page {doc.page}")

    doc.build(
        [
            Paragraph(
                escape(title), ParagraphStyle("title", fontName=font, fontSize=17, leading=22)
            ),
            Spacer(1, 8),
            cell(description),
            Spacer(1, 12),
            table,
        ],
        onFirstPage=footer,
        onLaterPages=footer,
    )
    return out.getvalue()


def export_bytes(fmt, title, headers, rows, description):
    if fmt == "pdf":
        return pdf_report(title, headers, rows, description), "application/pdf"
    clean = [[safe_cell(v) for v in row] for row in rows]
    if fmt == "csv":
        out = StringIO(newline="")
        writer = csv.writer(out)
        writer.writerow(headers)
        writer.writerows(clean)
        return ("\ufeff" + out.getvalue()).encode("utf-8"), "text/csv; charset=utf-8"
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.append([title])
    ws.append([description])
    ws.append(headers)
    for row in clean:
        ws.append(row)
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{max(3, ws.max_row)}"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 44
    ws["A1"].font = Font(size=17, bold=True, color="117D69")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    for c in ws[3]:
        c.fill = PatternFill("solid", fgColor="117D69")
        c.font = Font(bold=True, color="FFFFFF")
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 23
    for row in ws.iter_rows(min_row=4):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if c.row % 2 == 0:
                c.fill = PatternFill("solid", fgColor="EFF6F3")
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:3"
    out = BytesIO()
    wb.save(out)
    return out.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
