"""Bounded, strict imports. Parse before mutation; one transaction per entire file."""

import csv
import hashlib
import io
import zipfile
from xml.etree.ElementTree import ParseError

from defusedxml.common import DefusedXmlException
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import transaction
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from apps.accounts.policies import can
from apps.audit.services import record_event

from .models import Block, District, GramaPanchayat, LocationImport, State

COLUMNS = [
    "state_code",
    "state_name_en",
    "state_name_ml",
    "district_code",
    "district_name_en",
    "district_name_ml",
    "district_order",
    "block_code",
    "block_name_en",
    "block_name_ml",
    "block_lsg_code",
    "sec_local_body_code",
    "panchayat_name_en",
    "panchayat_name_ml",
    "panchayat_lsg_code",
    "official_email",
    "office_phone",
    "active",
]
REQUIRED = {
    "state_code",
    "state_name_en",
    "district_code",
    "district_name_en",
    "block_code",
    "block_name_en",
    "sec_local_body_code",
    "panchayat_name_en",
}
MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000


def parse_file(upload):
    data = upload.read(MAX_BYTES + 1)
    if len(data) > MAX_BYTES:
        raise ValidationError("The file exceeds the 5 MB limit.")
    name = str(upload.name).lower()
    try:
        if name.endswith(".csv"):
            reader = csv.reader(io.StringIO(data.decode("utf-8-sig")), strict=True)
            raw = []
            for i, row in enumerate(reader):
                if i > MAX_ROWS:
                    raise ValidationError("The file exceeds 5,000 data rows.")
                if len(row) > len(COLUMNS):
                    raise ValidationError("The file has too many columns.")
                raw.append(row)
        elif name.endswith(".xlsx"):
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if (
                    len(archive.infolist()) > 2000
                    or sum(item.file_size for item in archive.infolist()) > 25 * 1024 * 1024
                ):
                    raise ValidationError("The workbook expands beyond the safe size limit.")
            workbook = load_workbook(
                io.BytesIO(data), read_only=True, data_only=False, keep_links=False
            )
            try:
                if len(workbook.worksheets) != 1:
                    raise ValidationError("Use a workbook with exactly one sheet.")
                sheet = workbook.worksheets[0]
                if sheet.max_column and sheet.max_column > len(COLUMNS):
                    raise ValidationError("The workbook has too many columns.")
                raw = []
                for i, cells in enumerate(sheet.iter_rows()):
                    if i > MAX_ROWS:
                        raise ValidationError("The workbook exceeds 5,000 data rows.")
                    if any(cell.data_type == "f" for cell in cells):
                        raise ValidationError(
                            "Formulas are not accepted. Paste values before importing."
                        )
                    raw.append([cell.value for cell in cells])
            finally:
                workbook.close()
        else:
            raise ValidationError("Only UTF-8 CSV and XLSX files are supported.")
    except ValidationError:
        raise
    except (
        UnicodeError,
        csv.Error,
        zipfile.BadZipFile,
        KeyError,
        ValueError,
        OSError,
        DefusedXmlException,
        ParseError,
        InvalidFileException,
    ) as exc:
        raise ValidationError(
            "The file is malformed or is not a supported CSV/XLSX document."
        ) from exc
    if not raw:
        raise ValidationError("The file is empty.")
    headers = [str(value or "").strip() for value in raw[0]]
    if len(set(headers)) != len(headers):
        raise ValidationError("Duplicate column headers are not allowed.")
    if REQUIRED - set(headers) or set(headers) - set(COLUMNS):
        raise ValidationError(
            "Columns do not match the location template. Required: " + ", ".join(sorted(REQUIRED))
        )
    rows, codes, parents = [], set(), {}
    for number, raw_row in enumerate(raw[1:], start=2):
        if not any(value is not None and str(value).strip() for value in raw_row):
            continue
        if len(raw_row) != len(headers):
            raise ValidationError(f"Row {number}: column count does not match the header.")
        row = {key: "" for key in COLUMNS}
        row.update(
            {
                key: str(value).strip() if value is not None else ""
                for key, value in zip(headers, raw_row, strict=True)
            }
        )
        if any(not row[key] for key in REQUIRED):
            raise ValidationError(f"Row {number}: a required code or name is missing.")
        if any(len(value) > 2000 for value in row.values()):
            raise ValidationError(f"Row {number}: a cell exceeds the maximum length.")
        code = row["sec_local_body_code"]
        if code in codes:
            raise ValidationError(f"Row {number}: duplicate SEC Panchayat code {code}.")
        codes.add(code)
        try:
            row["district_order"] = int(row["district_order"] or "0")
        except ValueError as exc:
            raise ValidationError(f"Row {number}: district_order must be an integer.") from exc
        active = row["active"].lower()
        if active not in {"", "true", "false", "1", "0"}:
            raise ValidationError(f"Row {number}: active must be true or false.")
        row["active"] = active in {"", "true", "1"}
        # Repeated parents must have identical definitions, independent of row order.
        for kind, code_key, keys in [
            ("state", "state_code", ["state_name_en", "state_name_ml"]),
            (
                "district",
                "district_code",
                ["state_code", "district_name_en", "district_name_ml", "district_order"],
            ),
            (
                "block",
                "block_code",
                ["district_code", "block_name_en", "block_name_ml", "block_lsg_code"],
            ),
        ]:
            identity = (kind, row[code_key])
            definition = tuple(row[key] for key in keys)
            if identity in parents and parents[identity] != definition:
                raise ValidationError(
                    f"Row {number}: conflicting definitions for {kind} {row[code_key]}."
                )
            parents[identity] = definition
        row["_line"] = number
        rows.append(row)
    if not rows:
        raise ValidationError("At least one Panchayat row is required.")
    return rows, hashlib.sha256(data).hexdigest()


def _upsert(model, key, code, values, update_existing, summary):
    obj = model.objects.select_for_update().filter(**{key: code}).first()
    if obj is None:
        obj = model(**{key: code}, **values)
        obj.full_clean()
        obj.save()
        summary["created"][model._meta.model_name] += 1
    else:
        changed = {field: value for field, value in values.items() if getattr(obj, field) != value}
        if changed:
            if not update_existing:
                raise ValidationError(
                    f"{model.__name__} {code} differs from the master. Enable update_existing to change it."
                )
            old = {field: getattr(obj, field) for field in changed}
            for field, value in changed.items():
                setattr(obj, field, value)
            obj.revision += 1
            obj.full_clean()
            obj.save()
            summary["updated"][model._meta.model_name] += 1
            summary["changes"].append(
                {"model": model._meta.label_lower, "code": code, "old": old, "new": changed}
            )
        else:
            summary["unchanged"][model._meta.model_name] += 1
    return obj


@transaction.atomic
def import_locations(
    *, rows, sha256, source, actor=None, system=False, dry_run=False, update_existing=False
):
    if not system and not can(actor, "locations.manage"):
        raise PermissionDenied
    if not source.strip() or len(source) > 500:
        raise ValidationError("A source reference of at most 500 characters is required.")
    from apps.organization.models import OrganizationSetting

    # Shared singleton lock serializes import batches on production databases.
    OrganizationSetting.objects.select_for_update().get(pk=1)
    kinds = ["state", "district", "block", "gramapanchayat"]
    summary = {key: dict.fromkeys(kinds, 0) for key in ["created", "updated", "unchanged"]}
    summary["changes"] = []
    summary["rows"] = len(rows)
    seen = {}
    for row in rows:
        try:
            state_key = ("state", row["state_code"])
            if state_key not in seen:
                seen[state_key] = _upsert(
                    State,
                    "code",
                    row["state_code"],
                    {"name_en": row["state_name_en"], "name_ml": row["state_name_ml"]},
                    update_existing,
                    summary,
                )
            district_key = ("district", row["district_code"])
            if district_key not in seen:
                seen[district_key] = _upsert(
                    District,
                    "district_code",
                    row["district_code"],
                    {
                        "state_id": seen[state_key].pk,
                        "name_en": row["district_name_en"],
                        "name_ml": row["district_name_ml"],
                        "display_order": row["district_order"],
                    },
                    update_existing,
                    summary,
                )
            block_key = ("block", row["block_code"])
            if block_key not in seen:
                seen[block_key] = _upsert(
                    Block,
                    "block_code",
                    row["block_code"],
                    {
                        "district_id": seen[district_key].pk,
                        "name_en": row["block_name_en"],
                        "name_ml": row["block_name_ml"],
                        "lsg_code": row["block_lsg_code"],
                    },
                    update_existing,
                    summary,
                )
            _upsert(
                GramaPanchayat,
                "sec_local_body_code",
                row["sec_local_body_code"],
                {
                    "block_id": seen[block_key].pk,
                    "name_en": row["panchayat_name_en"],
                    "name_ml": row["panchayat_name_ml"],
                    "lsg_code": row["panchayat_lsg_code"],
                    "official_email": row["official_email"],
                    "office_phone": row["office_phone"],
                    "active": row["active"],
                },
                update_existing,
                summary,
            )
        except ValidationError as exc:
            raise ValidationError(f"Row {row['_line']}: {'; '.join(exc.messages)}") from exc
    if dry_run:
        transaction.set_rollback(True)
        return summary
    batch = LocationImport.objects.create(
        source=source, sha256=sha256, imported_by=actor, summary=summary
    )
    record_event(
        actor=actor,
        action="locations.imported",
        entity=batch,
        new_values={"source": source, "sha256": sha256, **summary},
    )
    return summary
