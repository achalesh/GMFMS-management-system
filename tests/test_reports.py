from datetime import timedelta
from io import BytesIO
from unittest.mock import patch

from django.test import Client, TestCase
from django.utils import timezone
from openpyxl import load_workbook
from pypdf import PdfReader

from apps.audit.models import AuditLog
from apps.facilitators.models import Facilitator
from apps.reports.data import report_data, snapshot, summarize
from apps.reports.exporting import export_bytes
from apps.reports.forms import ReportFilter
from tests import test_registry as fixtures
from tests.factories import assign


class ReportTests(TestCase):
    setUp = fixtures.RegistryTests.setUp
    make_app = fixtures.RegistryTests.make_app
    act = fixtures.RegistryTests.act
    begin = fixtures.RegistryTests.begin
    approve = fixtures.RegistryTests.approve
    change = fixtures.RegistryTests.change

    def stats(self, user=None, filters=None):
        return summarize(*snapshot(user or self.admin, filters or {}))

    def test_empty_and_pending_counts(self):
        m = self.stats()
        self.assertEqual(
            (m["panchayats"], m["active"], m["pending"], m["vacant"], m["coverage"]),
            (2, 0, 1, 2, 0),
        )
        self.assertEqual(self.stats(filters={"block": self.sibling})["coverage"], 0)

    def test_current_primary_coverage_and_assistant_does_not_fill_vacancy(self):
        f = self.approve()
        self.assertEqual(self.stats()["coverage"], 50)
        a = f.current_appointment
        a.role = "ASSISTANT"
        a.save()
        self.assertEqual((self.stats()["active"], self.stats()["covered"]), (1, 0))

    def test_future_ended_and_expired_appointments(self):
        f = self.approve()
        a = f.current_appointment
        a.effective_from = timezone.localdate() + timedelta(days=1)
        a.save()
        self.assertEqual(self.stats()["active"], 0)
        a.effective_from = timezone.localdate()
        a.effective_to = timezone.localdate() - timedelta(days=1)
        a.save()
        self.assertEqual(self.stats()["expired"], 1)
        a.effective_to = None
        a.ended_at = timezone.now()
        a.save()
        self.assertEqual(self.stats()["active"], 0)

    def test_suspension_revocation_and_expiry_reflected(self):
        f = self.approve()
        self.change(f, "suspend")
        self.assertEqual(self.stats()["suspended"], 1)
        self.assertEqual(self.stats()["covered"], 0)
        self.change(f, "reactivate")
        Facilitator.objects.filter(pk=f.pk).update(
            valid_until=timezone.localdate() - timedelta(days=1)
        )
        self.assertEqual(self.stats()["expired"], 1)

    def test_expiring_boundary_uses_appointment_end(self):
        f = self.approve()
        a = f.current_appointment
        a.effective_to = timezone.localdate() + timedelta(days=30)
        a.save()
        self.assertEqual(self.stats()["expiring"], 1)
        self.assertEqual(len(report_data(self.admin, {"report": "expiring"})[1]), 1)
        a.effective_to += timedelta(days=1)
        a.save()
        self.assertEqual(self.stats()["expiring"], 0)

    def test_inactive_location_excluded(self):
        self.approve()
        self.gp.active = False
        self.gp.save()
        self.assertEqual((self.stats()["panchayats"], self.stats()["active"]), (1, 0))

    def test_district_scope_and_denied_deep_link(self):
        self.approve()
        self.make_app(gp=self.other_gp, name="Other private applicant")
        self.assertEqual(self.stats(self.regional)["panchayats"], 1)
        self.client.force_login(self.regional)
        self.assertEqual(
            self.client.get(f"/dashboard/blocks/{self.other_block.pk}/").status_code, 404
        )
        self.assertNotContains(
            self.client.get("/reports/?report=applications"), "Other private applicant"
        )

    def test_mixed_role_export_scope_cannot_widen(self):
        assign(self.regional, "VIEWER", "STATE")
        self.client.force_login(self.regional)
        self.assertContains(self.client.get("/reports/"), "Other Panchayat")
        r = self.client.post("/reports/", {"format": "csv", "reason": "District planning"})
        self.assertNotIn(b"Other Panchayat", r.content)
        r = self.client.post(
            f"/reports/?district={self.other_block.district_id}",
            {"format": "csv", "reason": "District planning"},
        )
        self.assertEqual(r.status_code, 400)

    def test_viewer_contact_and_application_restrictions(self):
        self.approve()
        self.client.force_login(self.viewer)
        r = self.client.get("/reports/?report=facilitators")
        self.assertNotContains(r, "9876543210")
        self.assertContains(r, "Restricted")
        for kind in ["applications", "skills", "equipment", "progress"]:
            self.assertEqual(self.client.get("/reports/?report=" + kind).status_code, 403)
        self.assertEqual(
            self.client.post("/reports/", {"format": "csv", "reason": "No permission"}).status_code,
            403,
        )

    def test_hierarchy_dates_and_unknown_filters_rejected(self):
        for data in [
            {"district": self.district.pk, "block": self.other_block.pk},
            {"start": "2026-02-02", "end": "2026-01-01"},
            {"report": "facilitators", "status": "REJECTED"},
            {"report": "bogus"},
            {"report": "coverage", "role": "PRIMARY"},
        ]:
            self.assertFalse(ReportFilter(data, user=self.admin).is_valid())

    def test_registration_and_status_filters(self):
        self.approve()
        h, rows = report_data(
            self.admin,
            {
                "report": "facilitators",
                "start": timezone.localdate(),
                "end": timezone.localdate(),
                "status": "ACTIVE",
            },
        )
        self.assertEqual(len(rows), 1)
        self.assertEqual(
            report_data(
                self.admin,
                {"report": "facilitators", "start": timezone.localdate() + timedelta(days=1)},
            )[1],
            [],
        )

    def test_exports_audited_and_private_fields_absent(self):
        f = self.approve()
        self.client.force_login(self.admin)
        for fmt in ["csv", "xlsx", "pdf"]:
            r = self.client.post(
                "/reports/?report=facilitators",
                {"format": fmt, "reason": "Authorized district report"},
            )
            self.assertEqual(r.status_code, 200)
            self.assertIn("no-store", r["Cache-Control"])
            self.assertNotIn(f.verification_token.encode(), r.content)
            self.assertNotIn(b"Private home address", r.content)
            if fmt == "xlsx":
                ws = load_workbook(BytesIO(r.content)).active
                self.assertEqual(ws["A4"].value, f.facilitator_number)
                self.assertEqual(ws.freeze_panes, "A4")
            if fmt == "pdf":
                self.assertIn(
                    f.facilitator_number, PdfReader(BytesIO(r.content)).pages[0].extract_text()
                )
        self.assertEqual(AuditLog.objects.filter(action="reports.exported").count(), 3)

    def test_export_requires_reason_csrf_and_auth(self):
        self.assertEqual(self.client.get("/reports/").status_code, 302)
        self.client.force_login(self.admin)
        self.assertEqual(
            self.client.post("/reports/", {"format": "csv", "reason": "  "}).status_code, 400
        )
        c = Client(enforce_csrf_checks=True)
        c.force_login(self.admin)
        self.assertEqual(c.post("/reports/", {"format": "csv", "reason": "test"}).status_code, 403)
        self.assertFalse(AuditLog.objects.filter(action="reports.exported").exists())

    def test_formula_injection_and_unicode_export(self):
        values = ['=HYPERLINK("bad")', " +cmd", "@SUM(1)", "-1", "അചൽ കുമാർ"]
        data, _ = export_bytes("xlsx", "Test", ["Name"], [[v] for v in values], "QA")
        ws = load_workbook(BytesIO(data)).active
        self.assertTrue(all(ws.cell(i, 1).data_type != "f" for i in range(4, 9)))
        self.assertEqual(ws["A8"].value, values[-1])
        data, _ = export_bytes("csv", "Test", ["Name"], [[v] for v in values], "QA")
        self.assertIn("' +cmd", data.decode("utf-8-sig"))

    def test_appointment_history_and_skills_equipment_reports(self):
        f = self.approve()
        for kind in ["history", "skills", "equipment", "progress"]:
            h, rows = report_data(self.admin, {"report": kind})
            self.assertEqual(len(rows), 1)
        self.assertEqual(
            report_data(self.admin, {"report": "history"})[1][0][0], f.facilitator_number
        )

    def test_preview_pagination_and_no_audit_on_get(self):
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get("/reports/?page=9999").status_code, 200)
        self.assertFalse(AuditLog.objects.filter(action="reports.exported").exists())

    def test_dashboard_links_and_headers(self):
        self.approve()
        self.client.force_login(self.admin)
        for url in [
            "/",
            f"/dashboard/districts/{self.district.pk}/",
            f"/dashboard/blocks/{self.block.pk}/",
        ]:
            r = self.client.get(url)
            self.assertEqual(r.status_code, 200)
            self.assertContains(r, "Network coverage")
            self.assertIn("no-store", r["Cache-Control"])

    def test_export_audit_failure_returns_no_download(self):
        self.client.force_login(self.admin)
        with patch(
            "apps.reports.views.record_event", side_effect=RuntimeError("Audit unavailable")
        ):
            with self.assertRaises(RuntimeError):
                self.client.post("/reports/", {"format": "csv", "reason": "test"})

    def test_pdf_page_breaks_repeat_header(self):
        payload, _ = export_bytes(
            "pdf",
            "Test",
            ["Name", "Status"],
            [["Facilitator " + str(i), "ACTIVE"] for i in range(100)],
            "Synthetic QA",
        )
        pages = PdfReader(BytesIO(payload)).pages
        self.assertGreater(len(pages), 1)
        for page in pages:
            self.assertIn("Name", page.extract_text())
            self.assertIn("Page ", page.extract_text())

    def test_query_count_does_not_grow_with_facilitators(self):
        from django.db import connection
        from django.test.utils import CaptureQueriesContext

        self.approve()
        with CaptureQueriesContext(connection) as first:
            report_data(self.admin, {"report": "facilitators"})
        self.approve(self.make_app(gp=self.other_gp, name="Second facilitator"))
        with CaptureQueriesContext(connection) as second:
            report_data(self.admin, {"report": "facilitators"})
        self.assertEqual(len(first), len(second))

    def test_historical_appointment_scope_after_transfer(self):
        f = self.approve()
        self.change(f, "transfer", panchayat_id=self.other_gp.pk)
        rows = report_data(self.regional, {"report": "history"})[1]
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][4], self.gp.name_en)
        self.assertEqual(report_data(self.regional, {"report": "facilitators"})[1], [])

    def test_pdf_rejects_unsupported_glyphs(self):
        from django.core.exceptions import ValidationError

        with self.assertRaises(ValidationError):
            export_bytes("pdf", "Test", ["Name"], [["Unsupported emoji 😀"]], "QA")

    def test_pdf_limit_and_service_permission(self):
        from django.core.exceptions import PermissionDenied, ValidationError

        with self.assertRaises(ValidationError):
            export_bytes("pdf", "Test", ["Name"], [["QA"]] * 1501, "QA")
        with self.assertRaises(PermissionDenied):
            report_data(self.viewer, {}, export=True)

    def test_block_coordinator_scope(self):
        self.approve()
        self.assertEqual(self.stats(self.coordinator)["panchayats"], 1)
        self.client.force_login(self.coordinator)
        self.assertEqual(self.client.get(f"/dashboard/blocks/{self.sibling.pk}/").status_code, 404)
        self.assertEqual(
            self.client.post(
                "/reports/", {"format": "xlsx", "reason": "No export role"}
            ).status_code,
            403,
        )

    def test_pending_filter_combines_all_unresolved_states(self):
        self.make_app(gp=self.other_gp, name="Second pending")
        self.begin()
        form = ReportFilter({"report": "applications", "status": "PENDING"}, user=self.admin)
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(len(report_data(self.admin, form.cleaned_data)[1]), 2)
        self.assertEqual(
            len(report_data(self.admin, {"report": "applications", "status": "REJECTED"})[1]), 0
        )

    def test_multiple_primary_facilitators_cover_one_panchayat(self):
        self.approve()
        other = self.approve(self.make_app(gp=self.other_gp, name="Another primary"))
        a = other.current_appointment
        a.panchayat = self.gp
        a.save()
        m = self.stats()
        self.assertEqual((m["active"], m["covered"], m["vacant"]), (2, 1, 1))

    def test_contact_access_does_not_expand_with_viewer_role(self):
        self.approve(self.make_app(gp=self.other_gp, name="Other facilitator", mobile="9876543999"))
        assign(self.regional, "VIEWER", "STATE")
        rows = report_data(self.regional, {"report": "facilitators"})[1]
        self.assertEqual(rows[0][-1], "Restricted")
